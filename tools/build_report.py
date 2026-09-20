from pathlib import Path
import csv, json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT
OUT = PROJECT / 'docs' / 'Karolis_Lapinskas_Galutinio_darbo_ataskaita.docx'
source = (PROJECT/'docs/galutinis_darbas.md').read_text(encoding='utf-8')
metrics = json.loads((PROJECT/'results/final/final_test_metrics.json').read_text())
summary = list(csv.DictReader((PROJECT/'results/final/model_summary.csv').open(encoding='utf-8')))
doc = Document()
sec = doc.sections[0]
sec.page_width,sec.page_height = Inches(8.5),Inches(11)
sec.top_margin,sec.bottom_margin = Inches(.7),Inches(.65)
sec.left_margin,sec.right_margin = Inches(.8),Inches(.8)
sec.footer_distance = Inches(.3)
for style in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=doc.styles[style];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(11)
    s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=True
for el in doc.styles.element.iter():
    for attr in list(el.attrib):
        if attr.split('}')[-1] in ['themeColor','themeTint','themeShade','asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
            del el.attrib[attr]
    if el.tag==qn('w:pBdr'):
        el.getparent().remove(el)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.core_properties.title='Kredito paraiškų rizikos klasifikavimo galutinė ataskaita'
doc.core_properties.author='Karolis Lapinskas'
doc.core_properties.subject='PEPfm-26 Intelektualiosios sistemos'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Karolis Lapinskas  |  ');run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text,style=None):
    v=doc.add_paragraph(text,style);v.paragraph_format.widow_control=True
    return v
def h(text): return doc.add_heading(text,level=2)
def page(title):
    doc.add_page_break();doc.add_heading(title,level=1)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [6.9/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for cell,txt in zip(t.rows[0].cells,headers):cell.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge)
                for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+k),v)
                borders.append(b)
            tcpr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                n=OxmlElement('w:'+edge);n.set(qn('w:w'),'85');n.set(qn('w:type'),'dxa');mar.append(n)
            tcpr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DFE8F1' if ri==0 else 'FFFFFF');tcpr.append(sh)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(1);par.paragraph_format.line_spacing=1.04
                for r in par.runs:r.font.size=Pt(9.5);r.bold=(ri==0)
    p('') .paragraph_format.space_after=Pt(0)
    return t
def source_section(start,end):
    s=source.split(start,1)[1].split(end,1)[0].strip()
    return [x.strip() for x in s.split('\n\n') if x.strip()]
def paragraphs(blocks):
    for x in blocks:
        if not x.startswith(('|','!','#')):p(x)
def picture(path,width=6.5,caption=''):
    q=doc.add_paragraph();q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(PROJECT/path),width=Inches(width))
    if caption:p(caption,'Caption')
def mr(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
def box(tag,*children):
    n=OxmlElement('m:'+tag)
    for x in children:
        if not isinstance(x,str) and x.tag==qn('m:e') and tag in ['e','num','den']:
            for child in list(x):n.append(child)
        else:n.append(mr(x) if isinstance(x,str) else x)
    return n
def frac(a,b):return box('f',box('num',a),box('den',b))
def sub(a,b):return box('sSub',box('e',a),box('sub',b))
def sup(a,b):return box('sSup',box('e',a),box('sup',b))
def eq(*parts):
    par=doc.add_paragraph();par.paragraph_format.space_after=Pt(8)
    math=box('oMath',*parts);par._p.append(math)
    return par

doc.add_paragraph('Kredito paraiškų rizikos\nklasifikavimas','Title')
p('Galutinio darbo ataskaita','Subtitle')
p('Karolis Lapinskas\nPEPfm-26\nIntelektualiosios sistemos\n2026 m. rugsėjo 20 d.')
h('Darbo esmė')
p('Šiame darbe sukurtas ir patikrintas prototipas, kuris iš kredito paraiškos požymių įvertina blogos rizikos tikimybę. Palyginti šeši variantai, paaiškinta nevienoda dviejų klaidų kaina ir parodyta, kaip sprendimo slenkstis keičia rezultatą. Ataskaita nuosekliai aprašo duomenis, įgyvendinimą, bandymus ir jų ribas.')
h('Pagrindinis sprendimas')
p('Pasirinktas atraminių vektorių metodas SVM su RBF branduoliu ir sigmoidine tikimybių kalibracija. Jis pasiekė mažiausią vidutinę klaidų kainą plėtros duomenyse. Tai pagrindžia pasirinkimą šiame eksperimente, bet neįrodo, kad SVM visais atvejais yra geriausias kredito rizikos metodas.')
table(['Kas patikrinta','Gautas rezultatas'],[
('Duomenys','1 000 paraiškų, 20 požymių, 2 klasės'),
('Modelių palyginimas','6 metodai × 15 bendrų CV skaidymų'),
('Pasirinkto SVM vidutinė CV kaina','0,5583 sąlyginio vieneto paraiškai'),
('Galutinis testas','200 paraiškų; klaidų kaina 0,525'),
('Veikimo patikra','11 praėjusių testų ir naujų CSV įrašų prognozės')],[2.7,4.2])
p('Svarbiausias kompromisas: galutiniame teste aptikta 57 iš 60 blogos rizikos atvejų, bet blogai rizikai priskirta ir 90 iš 140 geros rizikos atvejų. Todėl darbe vertinama ne vien sėkmė, bet ir klaidingų atmetimų pasekmės.')
p('Kodas ir rezultatų failai: https://github.com/KarolisLLL/IS_EXAM')

page('1 Problema ir praktinė sprendimo vertė')
paragraphs(source_section('### 1.1 Problemos supratimas','### 1.2 Sprendimo etapai'))
h('Ką šiame darbe reiškia klasė')
p('„Good“ reiškia geros rizikos etiketę istoriniame duomenų rinkinyje, „bad“ – blogos rizikos etiketę. Tai nėra šiuolaikinio banko sprendimas ar konkretaus žmogaus būsimo mokėjimo garantija. Kai tekste vartojami žodžiai „patvirtinimas“ ir „atmetimas“, jie reiškia hipotetinę klasės pritaikymo pasekmę.')
table(['Tikra klasė','Prognozuota klasė','Prasmė','Kaina'],[
('good','good','Teisingai atpažinta gera rizika','0'),
('good','bad','FP – klaidingas geros rizikos atmetimas','1'),
('bad','good','FN – klaidingas blogos rizikos patvirtinimas','5'),
('bad','bad','Teisingai aptikta bloga rizika','0')],[1,1.25,3.9,.75])

page('2 Kas buvo atlikta ir kam to reikėjo')
table(['Etapas','Atliktas darbas','Įrodymas'],[
('Duomenų gavimas','Atsisiųstas credit-g ir patikrintos klasės bei požymiai','data_manifest.json'),
('Paruošimas','Užpildomos tuščios reikšmės, koduojamos kategorijos, keičiami skaičių masteliai','src/preprocess.py'),
('Modeliai','Įgyvendintos dvi pastovios taisyklės, logistinė regresija, SVM, MLP ir fuzzy','src/models.py; src/fuzzy.py'),
('Palyginimas','Visi metodai vertinti tuose pačiuose 15 išorinių skaidymų','outer_fold_results.csv'),
('Pasirinkimas','Modelis ir slenkstis parinkti iš plėtros duomenų','model_summary.csv'),
('Papildomi bandymai','Atlikta abliacija, trūkstamų reikšmių bei grupių diagnostika','Atskiri rezultatų CSV'),
('Naudojimas','Išsaugotas modelis, veikia naujo CSV prognozavimas ir HTML peržiūra','src/predict.py; index.html')],[1.1,4.05,1.75])
h('Uždavinių prasmė skirtingiems specialistams')
paragraphs(source_section('### 1.3 Uždavinių prasmė skirtingų sričių specialistams','## 2 Alternatyvos'))
p('Darbo eiga išskaidyta taip, kad kiekvienas etapas turėtų patikrinamą rezultatą. Tai leidžia atskirti duomenų kokybės klaidą nuo modelio klaidos ir nuo sąmoningai pasirinktos sprendimo politikos.')

page('3 Duomenys ir jų paruošimas')
p('Naudotas viešas OpenML credit-g rinkinys, data_id = 31. Jame yra 1 000 eilučių ir 20 požymių: 700 geros rizikos ir 300 blogos rizikos atvejų. Pirminį German Credit rinkinį aprašo UCI [5]. Tikslinė klasė užkoduota good = 0, bad = 1. Duomenų manifestas fiksuoja šaltinį, versiją, stulpelius ir SHA-256 kontrolinę sumą.')
table(['Požymių grupė','Pavyzdžiai ir prasmė'],[
('Skaitiniai požymiai','duration – mėnesiai; credit_amount – istorinio rinkinio piniginiai vienetai; age – metai; existing_credits – kreditų skaičius; taip pat installment_commitment, residence_since, num_dependents.'),
('Kategoriniai požymiai','Sąskaitos ir santaupų būsena, kredito istorija, tikslas, darbo kategorija, būstas, asmeninė būsena ir kitos kategorijos. Visi pavadinimai pateikti examples/applications.csv.')],[1.6,5.3])
h('Kaip požymiai paverčiami modelio įvestimi')
p('Skaitinės trūkstamos reikšmės pakeičiamos mokymo dalies mediana. Tada skaitiniai požymiai standartizuojami. Pavyzdžiui, jei mokymo dalyje trukmės vidurkis būtų 24 mėnesiai, o standartinis nuokrypis 12, tai 36 mėnesių reikšmė po transformacijos būtų 1. Šie skaičiai yra tik paaiškinantis pavyzdys, ne rinkinio statistika.')
p('Kategorijų negalima tiesiog sunumeruoti ir laikyti didėjančiais dydžiais. Todėl naudojamas atskirų indikatorių kodavimas, vadinamas one-hot. Supaprastintame būsto pavyzdyje „own“ galėtų būti (1, 0, 0), „rent“ – (0, 1, 0), „for free“ – (0, 0, 1). Nenurodoma, kad viena būsto kategorija yra tris kartus didesnė už kitą.')
p('Kategorinės trūkstamos reikšmės pakeičiamos dažniausia mokymo dalies kategorija. Nežinoma nauja kategorija gauna nulinį savo požymio indikatorių vektorių. Tai leidžia programai veikti, tačiau neįrodo, kad jos prognozė tokiam nepažįstamam atvejui bus patikima.')
h('Kaip išvengta duomenų nutekėjimo')
p('Medianos, vidurkiai, nuokrypiai ir kategorijų sąrašai nustatomi tik mokymo dalyje. Validavimo ir testo eilutės transformuojamos jau išmoktomis taisyklėmis. Jei šios reikšmės būtų skaičiuojamos iš viso rinkinio, vertinimo duomenys iš anksto paveiktų modelį ir rezultatas galėtų būti pernelyg optimistinis. SVM atveju ta pati taisyklė taikoma ir kalibravimo skaidymams.')

page('4 Kaip veikia palyginti metodai')
h('Pastovios atskaitos taisyklės')
p('Daugumos taisyklė visada prognozuoja good. Ji pasiekia 70 % bendrą tikslumą, bet praleidžia visus bad atvejus. Esant 5:1 kainai, jos vidutinė kaina yra 0,30 × 5 = 1,50. Antroji taisyklė visada prognozuoja bad: ji nepalieka nepastebėtų bad atvejų, bet klaidingai atmeta visus good; kaina 0,70 × 1 = 0,70. Abi taisyklės padeda įvertinti, ar sudėtingesnis modelis iš tiesų duoda naudos.')
h('Logistinė regresija')
p('Modelis sudeda transformuotus požymius su išmoktais svoriais, o gautą skaičių paverčia tikimybe. Tai paprastas mokomas atskaitos metodas. Jis gerai tinka nedidelėms lentelinėms imtims ir leidžia nagrinėti koeficientų kryptį, tačiau sudėtingos netiesinės sąveikos savaime neatsiranda. Koeficientas taip pat nėra priežastinio poveikio įrodymas.')
h('Atraminių vektorių metodas SVM')
p('SVM ieško klases skiriančios ribos. RBF branduolys leidžia šiai ribai būti netiesinei: dvi paraiškos lyginamos pagal visų transformuotų požymių panašumą. Reguliavimo parametras C valdo mokymo klaidų ir ribos paprastumo kompromisą, o γ – panašumo vietiškumą. SVM balas papildomai kalibruojamas, kad būtų galima taikyti tikimybinį slenkstį.')
h('Daugiasluoksnis perceptronas MLP')
p('MLP yra neuroninis tinklas: įvesties požymiai perduodami vienam ar dviem paslėptiems sluoksniams, kurie mokosi netiesinių jų derinių. Darbe išbandyti nedideli tinklai, reguliavimas ir mokymosi greičiai. Taikytas ankstyvas stabdymas, iteracijų riba 400. Tinklas gali išmokti sudėtingą ryšį, bet nedidelėje imtyje jo rezultatą labiau veikia mokymo eiga ir parametrų pasirinkimas.')
h('Fuzzy tipo taisyklių sistema')
p('Fuzzy metodas riziką didina palaipsniui: suma gali būti „iš dalies didelė“, trukmė – „iš dalies ilga“. Šių narystės reikšmių svertinė suma sujungiama su sąskaitos ir santaupų kategorijų rizikos balais. Tai lengvai išaiškinama demonstracija. Svoriai pasirinkti rankiniu būdu; jie neišmokti iš duomenų ir nepatvirtinti ekspertų. Naudojama supaprastinta sistema, o ne pilna Mamdani ar Sugeno ekspertinė sistema.')

page('5 Metodų privalumai trūkumai ir alternatyvos')
table(['Metodas','Pliusai','Minusai'],[
('Logistinė regresija','Paprasta, greita, reguliuojama; patogi atskaita ir koeficientų analizė.','Tiesinis balas; sudėtingoms sąveikoms reikia papildomų požymių.'),
('RBF SVM','Netiesinė riba be rankinio visų sąveikų kūrimo; šiame bandyme mažiausia kaina.','Reikia mastelių keitimo, parametrų paieškos ir kalibravimo; sunkiau paaiškinti vieną sprendimą.'),
('MLP','Lankstus požymių sąveikų mokymasis; galima keisti tinklo dydį.','Mažoje imtyje didesnė persimokymo rizika; rezultatai priklauso nuo mokymo nustatymų.'),
('Fuzzy','Aiškios formulės ir palaipsnės taisyklės; patogu parodyti kodo ir matematikos ryšį.','Rankiniai svoriai, tik keturi naudojami požymiai; balas nėra kalibruota tikimybė.'),
('Pastovios taisyklės','Nereikia mokymo; leidžia patikrinti, ar sudėtingas modelis naudingas.','Neskiria individualių paraiškų; viena klaidų rūšis išlieka labai dažna.')],[1.15,2.75,3])
h('Kokie kiti metodai dar galėtų būti taikomi')
p('Sprendimų medis galėtų pateikti nuoseklias „jeigu – tada“ taisykles. Jo privalumas – lengvai peržiūrima sprendimo eiga, ypač kai medis negilus. Trūkumas – vienas gilus medis gali persimokyti ir pasikeisti dėl nedidelio mokymo duomenų pokyčio. Šiame darbe medis nebuvo išbandytas, todėl jo kokybė nežinoma.')
p('Atsitiktinis miškas sujungtų daug skirtingų medžių. Tai galimas būdas aprašyti netiesines sąveikas ir sumažinti vieno medžio nestabilumą [8]. Trūkumai – sudėtingesnis individualus paaiškinimas ir poreikis atskirai patikrinti tikimybių kalibraciją. Jo pranašumo prieš SVM šio darbo rezultatai neįrodo.')
p('Gradientinis medžių stiprinimas nuosekliai prideda modelius, koreguojančius ankstesnio junginio klaidas [9]. Jis yra pagrįsta tolesnio eksperimento alternatyva, tačiau reikalauja papildomo reguliavimo ir sąžiningo parametrų paieškos biudžeto. Pridėti jį po testo peržiūros ir iš karto paskelbti nauju laimėtoju būtų nekorektiška – reikėtų naujo atrankos bei patikros ciklo.')

page('6 Kodėl pasirinkti būtent šie metodai')
p('Metodai parinkti taip, kad būtų palygintos skirtingos sprendimo idėjos: paprasta tiesinė tikimybės formulė, netiesinė atraminių vektorių riba, neuroninis tinklas ir aiškinamos narystės taisyklės. Duomenų atsisiuntimui bei schemos tikrinimui intelektualaus metodo nereikia; tai įprasti deterministiniai programavimo veiksmai.')
h('Tiesioginis kredito rizikos pagrindimas')
p('Lessmann ir kt. (2015) atliko kredito rizikos klasifikavimo algoritmų palyginimą, įtraukdamas German Credit duomenis [1]. Tai tiesiogiai pagrindžia kelių modelių palyginimą šiame dalykiniame uždavinyje ir logistinės regresijos naudojimą kaip atskaitą. Straipsnis nereiškia, kad konkretus metodas privalo laimėti šio darbo skaidyme.')
h('SVM teorinis pagrindas')
p('Cortes ir Vapnik (1995) aprašo atraminių vektorių klasifikavimą [2]. Šiame darbe RBF SVM pasirinktas kaip pagrindinis kandidatas todėl, kad rizika gali priklausyti nuo kelių požymių sąveikų: pavyzdžiui, didelė suma ir ilga trukmė kartu gali turėti kitokią reikšmę nei kiekvienas požymis atskirai. Tai patikrinama hipotezė, o ne iš anksto žinomas rezultatas.')
h('MLP ir fuzzy pagrindimas')
p('Cybenko (1989) pateikia teorinį sigmoidinių neuroninių tinklų aproksimavimo rezultatą [3]. Jis motyvuoja neuroninių modelių lankstumą, bet negarantuoja, kad šio projekto ReLU MLP gerai išmoks iš 800 plėtros įrašų. Tinkamumas sprendžiamas pagal eksperimento metrikas.')
p('Zadeh (1965) apibrėžia fuzzy aibes ir narystės laipsnius [4]. Ši idėja tinka palaipsniam „didelės sumos“ ar „ilgos trukmės“ aprašymui. Šaltinis nepagrindžia šiame kode pasirinktų konkrečių koeficientų. Todėl fuzzy modelis vertinamas kaip aiškinamas demonstracinis palyginimas.')
h('Ką reiškia pagrįstas pasirinkimas')
p('Literatūra paaiškina, kodėl metodą verta bandyti. Tik vienodomis sąlygomis atliktas eksperimentas parodo, kaip jis pasirodė šiame rinkinyje. Šių dviejų argumentų nereikia suplakti: teorinis lankstumas savaime nėra mažesnės klaidų kainos įrodymas.')

page('7 Mokymo ir vertinimo eiga')
paragraphs(source_section('### 3.1 Vertinimo protokolas ir pasirinkimo kriterijus','### 3.2 Rezultatai'))
h('Kryžminis vertinimas paprastai')
p('800 plėtros eilučių padalijamos į penkias dalis. Viename išoriniame žingsnyje apie 640 eilučių skiriama mokymui, 160 – vertinimui. Procesas kartojamas taip, kad kiekviena dalis pabūtų vertinimo dalimi. Visa penkių dalių procedūra pakartojama tris kartus. Stratifikavimas išlaiko panašų good ir bad santykį.')
p('Vidinis CV vyksta tik to išorinio žingsnio mokymo dalyje. Jis padeda rinktis parametrus ir gauti prognozes slenksčiui parinkti. Out-of-fold prognozė reiškia, kad konkretaus įrašo prognozę pateikė modelis, kuris to įrašo savo mokymo dalyje neturėjo. Išorinė vertinimo dalis skirta patikrinti jau pasirinktą sprendimą.')
p('200 galutinio testo eilučių paliekamos nuošalyje iki pasirinkimo pabaigos. Tai padeda atskirti mokymo ir parametrų paieškos sėkmę nuo veikimo atskiruose duomenyse. Šis skaidymas jau buvo naudotas ankstesnėje projekto versijoje, todėl pakartotas vykdymas nėra naujų nepriklausomų duomenų tyrimas.')

page('8 Kaip skaityti kokybės rodiklius')
table(['Rodiklis','Ką parodo','Kaip vertinti'],[
('Bendras tikslumas','Visų teisingų prognozių dalį.','Gali slėpti retą bad klasę; vien jo nepakanka.'),
('Balanced accuracy','Bad aptikimo ir good atpažinimo dalių vidurkį.','Abi klasės vienodai svarbios; didesnė reikšmė geresnė.'),
('AP','Kaip gerai bad atvejai surikiuojami aukščiau pagal riziką.','Didesnė reikšmė geresnė; čia tai average precision.'),
('Brier balas','Tikimybių kvadratinę paklaidą tikros 0 arba 1 klasės atžvilgiu.','Mažesnis geresnis; vertina tikimybės įverčio kokybę.'),
('Klaidų kaina','FP ir FN skaičius, įvertintus kainomis 1 ir 5.','Mažesnė geresnė; tai pagrindinis atrankos kriterijus.'),
('Kalibracija','Ar prognozuotos tikimybės atitinka stebėtą bad dažnį.','Geras atitikimas artėja prie grafiko įstrižainės.')],[1.35,2.75,2.8])
p('Failuose vartojamas stulpelio pavadinimas pr_auc, tačiau jo reikšmė skaičiuojama average_precision_score funkcija. Ataskaitoje ji žymima AP. Tai nėra trapeciniu būdu apskaičiuotas PR kreivės plotas.')
h('Kodėl rodikliai gali prieštarauti vienas kitam')
p('Mažesnis slenkstis paprastai padeda aptikti daugiau bad atvejų, bet padidina good atmetimą. Dėl to klaidų kaina gali sumažėti, nors bendras tikslumas pablogėja. Vienas modelis taip pat gali gerai surikiuoti riziką, tačiau pateikti pernelyg dideles tikimybes. Todėl pateikiami keli rodikliai, o laimėtojas renkamas pagal iš anksto apibrėžtą kriterijų.')
h('Kas matuojama sąlyginiais vienetais')
p('Kaina 0,525 reiškia vidutinę klaidų kainą vienai paraiškai, kai FP kainuoja 1, o FN – 5 vienetus. Tai nėra 0,525 euro ir nėra 52,5 % tikimybė. Tikram finansiniam įvertinimui reikėtų banko nuostolių, palūkanų, paskolų dydžių ir kitų veiklos duomenų.')

page('9 Modelių palyginimo rezultatai')
p('Lentelė sudaryta iš results/final/model_summary.csv. Pateikti 15 bendrų išorinių skaidymų vidurkiai. SD rodo kainos sklaidą tarp skaidymų; tai nėra pasikliautinasis intervalas.')
names={'svm':'SVM','mlp':'MLP','logistic':'Logistinė regresija','fuzzy':'Fuzzy','reject_all':'Visada bad','majority':'Visada good'}
def fmt(v,n=4):return f'{float(v):.{n}f}'.replace('.',',')
table(['Metodas','BA','AP','Brier','Kaina ± SD'],[[names[r['model']],fmt(r['balanced_accuracy_mean']),fmt(r['pr_auc_mean']),fmt(r['brier_score_mean']),fmt(r['expected_cost_mean'])+' ± '+fmt(r['expected_cost_std'])] for r in summary],[1.5,1.03,1.03,1.03,2.31])
picture(Path('results/final/model_comparison.png'),6.45,'1 pav. Metodų palyginimas pagal vidutinę klaidų kainą ir AP. Mažesnė kaina bei didesnis AP yra geresni.')
p('SVM kainos vidurkis mažiausias. Logistinė regresija ir SVM turi beveik tą patį AP, todėl pagal įrašų surikiavimą aiškaus SVM pranašumo nėra. Abu mokomi atskaitos ir intelektualieji modeliai kainos požiūriu lenkia visada bad bei visada good taisykles; fuzzy taip pat pagerina pastovias atskaitas, tačiau atsilieka nuo SVM.')

page('10 Pasirinktas SVM ir pasirinkimo argumentai')
blocks=source_section('### 3.2 Rezultatai ir pasirinkto metodo privalumai','## 4 Įvestis')
paragraphs([b for b in blocks if b.startswith(('Pasirinktas','SVM išskirtinis'))])
h('Koks SVM variantas galutinai išmokytas')
table(['Nustatymas','Reikšmė','Prasmė'],[
('Branduolys','RBF','Leidžia netiesinę klasifikavimo ribą.'),
('C','0,215443469','Mokymo klaidų ir reguliavimo kompromisas.'),
('γ','0,05','Nustato, kaip greitai mažėja panašumas tarp įrašų.'),
('Klasių svoriai','balanced','Mokant atsižvelgiama į klasių dažnių skirtumą.'),
('Kalibracija','Sigmoidinė, 3 dalys','SVM balas paverčiamas bad tikimybės įverčiu.'),
('Sprendimo slenkstis','0,15','bad prognozuojama, kai p_bad ≥ 0,15.')],[1.45,1.75,3.7])
p('Klasių svoriai ir sprendimo slenkstis nėra tas pats. Svoriai veikia modelio mokymą; slenkstis taikomas jau gautai tikimybei. SVM pasirinktas pagal CV kainą, o ne todėl, kad jo pavadinimas sudėtingesnis ar kad jis būtų geriausias pagal kiekvieną rodiklį.')

page('11 SVM formulės ir ryšys su programa')
h('Požymių standartizavimas')
eq(sub('z','j'),' = ',frac(box('e',sub('x','j'),' − ',sub('μ','j')),sub('σ','j')))
p('xⱼ – įvesties reikšmė po trūkstamų reikšmių užpildymo; μⱼ ir σⱼ – mokymo dalies vidurkis bei standartinis nuokrypis. Nulinės dispersijos atveju mastelis laikomas 1. Kodo vieta: src/preprocess.py, make_preprocessor.')
h('RBF branduolys ir SVM balas')
eq('K(z, ',sub('z','i'),') = exp(−γ ',sup(box('e','‖z − ',sub('z','i'),'‖'),'2'),')')
eq('f(z) = ',sub('Σ','i'),sub('α','i'),sub('y','i'),' K(z, ',sub('z','i'),') + b')
p('z – naujos paraiškos transformuotas vektorius, zᵢ – atraminis mokymo vektorius. αᵢ ir b išmokstami; γ parenkamas parametrų paieškoje. Šioje SVM formulėje yᵢ ∈ {−1, +1}, nors projekto duomenų lentelėse klasės koduojamos 0 ir 1. Kodo vieta: src/models.py, SVC(kernel="rbf").')
h('Balas paverčiamas tikimybe')
eq(sub('p','bad'),'(z) = ',frac('1','1 + exp(A f(z) + B)'))
p('A ir B nustatomi kalibruojant iš mokymo dalies skaidymų prognozių. Įgyvendinimas: CalibratedClassifierCV su method="sigmoid". Visa transformacijų grandinė yra kalibratoriaus viduje, kad jo vertinimo eilutės nepaveiktų paruošimo [7].')
h('Tikimybė paverčiama klase')
eq('ŷ = I(',sub('p','bad'),' ≥ t)')
eq('C(t) = ',frac(box('e',sub('c','FP'),' FP(t) + ',sub('c','FN'),' FN(t)'),'n'))
p('I yra indikatorius: sąlygai tenkinantis grąžina 1, kitu atveju 0. n – įrašų skaičius; c_FP = 1 ir c_FN = 5. t parenkamas iš 0,01–0,99 tinklelio pagal mažiausią mokymo out-of-fold kainą. Kodo vietos: choose_cost_threshold ir expected_cost faile src/evaluation.py; prognozavimo taisyklė – src/predict.py.')

page('12 Fuzzy taisyklė ir metrikų formulės')
h('Palaipsnė narystė')
eq('μ(v; l, h) = min(1, max(0, ',frac('v − l','h − l'),'))')
p('v – trukmė arba suma, l ir h – pasirinktos apatinė ir viršutinė ribos. Trukmei naudojama 6 ir 48 mėnesiai, sumai 1 000 ir 8 000. Jei trukmė 27 mėnesiai, narystė lygi (27−6)/(48−6) = 0,5. Trūkstama reikšmė fuzzy variante keičiama intervalo viduriu.')
eq('r = clip(0,08 + 0,26 ',sub('μ','trukmė'),' + 0,29 ',sub('μ','suma'),
       ' + 0,27 ',sub('r','sąskaita'),' + 0,18 ',sub('r','santaupos'),', 0,01, 0,99)')
p('Kategorijų rizikos balai paimami iš src/fuzzy.py žodynų. Pavyzdžiui, checking_status = „<0“ gauna 1,0, o „>=200“ – 0,25. Nežinomai kategorijai taikoma 0,55. Clip apriboja rezultatą. Formulė tiesiogiai įgyvendinta _linear_membership, _category_risk ir predict_proba funkcijose. Gautas r yra rankinis rizikos balas, ne patvirtinta kalibruota tikimybė.')
h('Subalansuotas tikslumas')
eq('BA = ',frac('1','2'),' (',frac('TP','TP + FN'),' + ',frac('TN','TN + FP'),')')
p('Pirmoji dalis rodo bad aptikimą, antroji – good atpažinimą. Taip abiem klasėms suteikiamas vienodas svoris.')
h('Tikimybių paklaida ir įrašų surikiavimas')
eq('Brier = ',frac('1','n'),sub('Σ','i'),sup(box('e','(',sub('p','i'),' − ',sub('y','i'),')'),'2'))
eq('AP = ',sub('Σ','k'),' (',sub('R','k'),' − ',sub('R','k−1'),') ',sub('P','k'))
p('Brier formulėje pᵢ yra prognozuota bad tikimybė, yᵢ – tikra 0 arba 1 klasė. AP formulėje Pₖ – tikslumas tarp pažymėtų bad prognozių, Rₖ – aptiktų tikrų bad dalis surikiuoto sąrašo taške. Visos metrikos įgyvendintos src/evaluation.py, metric_row.')

page('13 Galutinio testo rezultato interpretacija')
p('Atrankai nenaudotame 200 įrašų teste pasirinktas SVM su 0,15 slenksčiu gavo tokią klaidų matricą. Teigiama klasė šiame darbe yra bad.')
table(['Tikra klasė','Prognozė good','Prognozė bad','Iš viso'],[
('good','TN = 50','FP = 90','140'),('bad','FN = 3','TP = 57','60'),('Iš viso','53','147','200')],[1.6,1.9,1.9,1.5])
p('Iš 60 tikrų bad atvejų aptikti 57: bad aptikimas 95 %. Tačiau 90 iš 140 tikrų good atvejų klaidingai priskirti bad: klaidingų atmetimų dalis 64,3 %. Bendras tikslumas yra tik (50+57)/200 = 53,5 %. Tai ne programos veikimo klaida, o pasirinktos 5:1 kainos ir žemo slenksčio kompromisas.')
p('Galutinės metrikos: balanced accuracy 0,6536; AP 0,6784; Brier 0,1528. Kaina lygi (1 × 90 + 5 × 3) / 200 = 0,525. Šis skaičius mažesnis už visada bad taisyklės kainą 0,700, tačiau klaidingų atmetimų kiekis vis tiek didelis.')
h('Kas nutinka pakeitus slenkstį')
table(['Slenkstis','FP','FN','Bad aptikimas','Bendras tikslumas','Kaina'],[
('0,15','90','3','95,0 %','53,5 %','0,525'),
('0,50','19','24','60,0 %','78,5 %','0,695')],[.8,.6,.6,1.6,1.8,1.5])
p('Esant 0,50 slenksčiui bendras tikslumas didesnis, bet praleidžiami 24 bad atvejai vietoje 3. Kadangi kiekvienas toks praleidimas kainuoja penkis kartus daugiau, bendra klaidų kaina pakyla. Ši diagnostika parodo, kodėl negalima aklai rinktis 0,50 ar vien didžiausio bendro tikslumo.')
p('Slenksčių lentelė apskaičiuota tam pačiam galutiniam testui ir skirta paaiškinti poveikį. Pagal šią lentelę modelis ar slenkstis neperrenkami. Galutinį 0,15 slenkstį parinko plėtros duomenų procedūra.')
p('Jei tikimybės būtų tiksliai kalibruotos, teorinis slenkstis pagal šią kaštų matricą būtų 1/(1+5) ≈ 0,167. Empirinis 0,15 nuo jo gali skirtis dėl imties, kalibracijos paklaidos ir pasirinkto kandidatų tinklelio.')

page('14 Kalibracija abliacija ir atsparumas')
picture(Path('results/final/calibration_final_test.png'),5.0,'2 pav. Prognozuotos bad tikimybės ir stebėto bad dažnio palyginimas galutiniame teste.')
p('Grafiko įstrižainė reiškia idealų atitikimą: pavyzdžiui, tarp panašaus 0,20 įverčio paraiškų apie 20 % būtų bad. Kreivė nuo įstrižainės vietomis nukrypsta. Ties maždaug 0,65 prognoze stebėta bad dalis apie 0,43, o ties 0,55 – apie 0,73. Taigi kalibravimas nėra tobulo tikimybės tikslumo garantija; mažas testas riboja interpretaciją.')
table(['Bandymas','BA','AP','Brier','Kaina'],[
('Pagrindinis modelis','0,654','0,678','0,153','0,525'),
('Be age ir personal_status','0,674','0,650','0,160','0,510'),
('5 % trūkstamų reikšmių','0,657','0,662','0,155','0,520'),
('10 % trūkstamų reikšmių','0,663','0,658','0,154','0,525')],[2.7,1.05,1.05,1.05,1.05])
p('Abliacija reiškia pasirinktų požymių pašalinimą ir modelio mokymą iš naujo. Be amžiaus ir asmeninės būsenos kaina šiame teste sumažėjo, bet AP ir Brier pablogėjo. Tai nėra visapusiškas pagerėjimas ar priežastinio poveikio įrodymas. Šis variantas pagal testo rezultatą neperrenkamas nauju laimėtoju.')
p('Atsparumui tikrinti atsitiktinai maskuoti testo langeliai su 5 % ir 10 % tikimybe, naudojant fiksuotas sėklas. Modelis veikė ir šiame bandyme kaina beveik nepasikeitė, tačiau kiekvienai daliai atliktas tik vienas maskavimo bandymas. Tai netikrina sisteminio trūkumo ar ekonominių pokyčių; faktinis maskuotų langelių procentas gali nežymiai skirtis.')

page('15 Klaidų pavyzdžiai ir grupių diagnostika')
table(['Rinkinio eilutė','Tikra klasė','Prognozė','Bad įvertis','Paaiškinimas'],[
('409','bad','good','0,123','Įvertis mažesnis už 0,15: praleista bloga rizika.'),
('190','bad','good','0,142','Netoli slenksčio, bet vis tiek FN klaida.'),
('278','bad','good','0,147','Taip pat vos žemiau slenksčio.'),
('286','good','bad','0,690','Didelis rizikos įvertis, bet tikra etiketė good.')],[.9,.9,.9,1,3.2])
p('Eilutės numeris yra viešo rinkinio indeksas, ne žmogaus tapatybė. Programa išsaugo iki penkių kiekvieno klaidos tipo pavyzdžių. FN atveju imami mažiausi bad įverčiai, FP atveju – didžiausi. Matomi atvejai parodo, kad modelis gali klysti ir prie slenksčio, ir būdamas gana užtikrintas. Vien šie pavyzdžiai nepasako priežasties, kodėl paraiška buvo neteisingai įvertinta.')
h('Ar klaidos skirtingose grupėse vienodos')
table(['Grupė','Testo įrašai','Bad aptikimas','Good klaidingas atmetimas'],[
('Iki 25 metų','26','100,0 %','84,6 %'),('Nuo 25 metų','174','93,6 %','62,2 %'),
('female pagal personal_status','63','95,5 %','73,2 %'),('Kita personal_status grupė','137','94,7 %','60,6 %')],[2.5,1,1.45,1.95])
p('Skirtumai verti dėmesio: jaunesnėje grupėje klaidingų geros rizikos atmetimų dalis didesnė. Tačiau joje iš viso tik 26 įrašai, o atskirų klasių dar mažiau. Nepateikti pasikliautinieji intervalai, todėl negalima skelbti patikimai įrodyto grupių skirtumo ar jo priežasties.')
p('Istorinė personal_status kategorija sujungia kelias savybes. Iš jos išvesta female grupė nėra išsamus šiuolaikinės populiacijos aprašymas. Auditas yra aprašomoji diagnostika, ne teisingumo sertifikatas. Pašalinti du jautresnius požymius taip pat nereiškia automatiškai pašalinti visą informaciją apie grupes, nes ją gali netiesiogiai perteikti kiti požymiai.')

page('16 Programos paleidimas ir patikra')
p('Projektas parengtas GitHub saugykloje KarolisLLL/IS_EXAM. Paleidimo instrukcija yra README.md, priklausomybių versijos – requirements.txt, tiksli eksperimento aplinka ir nustatymai – results/final/run_manifest.json. Rekomenduojama Python 3.12. Visos toliau pateiktos komandos vykdomos iš projekto šakninio katalogo.')
for title,command,explanation in [
('Aplinkos paruošimas','python -m pip install -r requirements.txt','Prieš tai sukuriama ir aktyvinama virtualioji Python aplinka, kaip nurodyta README.'),
('Pilnas eksperimentas','python -m src.main --config config/experiment.yaml','Sukuria pilnus palyginimo, testo, abliacijos, atsparumo ir peržiūros rezultatus results/final kataloge.'),
('Greitas veikimo bandymas','python -m src.main --config config/experiment.yaml --quick','Mažesnis techninės grandinės patikrinimas. Rašo į results/quick, todėl nepakeičia pilno eksperimento rezultatų.'),
('Naujos paraiškos prognozė','python -m src.predict --input examples/applications.csv','CSV turi turėti visus 20 požymių stulpelių. Tikros klasės prognozuojant nereikia.'),
('Automatiniai testai','python -m pytest -q','Praėjo 11 vietinių testų. Pilnas eksperimentas ir naujų paraiškų komanda taip pat paleisti.')]:
    h(title);q=p(command);q.runs[0].font.name='Consolas';q.runs[0].font.size=Pt(9.5);p(explanation)
h('Ką parodo demonstracinės naujos paraiškos')
p('Dvi ranka sukurtos eilutės gavo bad įverčius apie 0,040 ir 0,758. Taikant 0,15 slenkstį, pirmoji priskirta good, antroji bad. Tai patvirtina, kad išsaugotas modelis gali priimti naują CSV. Kadangi tikros baigtys nežinomos, šis pavyzdys neįrodo prognozių tikslumo.')
p('Išvesties CSV pateikia row – eilutės numerį, bad_probability – rizikos įvertį, threshold – taikytą slenkstį ir predicted_class – good arba bad. Trūkstant viso privalomo požymio stulpelio arba pateikus tekstą skaitinio požymio vietoje grąžinama klaida. Tušti langeliai užpildomi išsaugoto modelio taisyklėmis.')
p('Rezultatus galima peržiūrėti atvėrus results/final/index.html. Jame yra grafikai, lentelės ir slenksčio valdiklis. Kodo testai tikrina duomenų klases, klaidų kainą, fuzzy ribas, slenksčio taikymą, netaisyklingą įvestį ir kalibravimo grandinės sandarą.')

page('17 Darbo ribos AI naudojimas ir išvados')
h('Ko šis eksperimentas neįrodo')
p('Duomenų rinkinys mažas ir istorinis. Nebuvo tikrinamas modelio veikimas kitu laikotarpiu ar kitoje šalyje, o 5:1 kaštai yra demonstracinė prielaida. Parametrų paieška ribota penkiais bandiniais, todėl blogesnis MLP ar kito metodo rezultatas gali priklausyti ir nuo pasirinkto skaičiavimo biudžeto. Kartotinio CV skaidymai persidengia; jų sklaida nėra nepriklausomų eksperimentų statistinio patvirtinimo pakaitalas.')
p('Galutinis testas nebuvo naudojamas šiame kode modelio šeimai, parametrams ar slenksčiui parinkti. Vis dėlto tas pats fiksuotas testinis skaidymas buvo peržiūrėtas ankstesnėje projekto versijoje. Todėl pateikiamas atkuriamas prototipo patikrinimas, o ne naujas nepriklausomas bankinės sistemos patvirtinimo tyrimas. Praktiniam taikymui reikėtų naujų duomenų ir tikrų veiklos kaštų.')
h('Kaip naudotas ir tikrintas AI')
p('AI padėjo rengti kodą, tekstą ir patikrų planą. Šaltiniai tikrinti pagal pirmines publikacijas ir oficialią dokumentaciją; skaičiai paimti iš vykdytos programos failų. AI naudojimo žurnalas saugomas docs/ai_usage_log.md.')
p('Patikros metu pataisyta SVM kalibravimo grandinė: duomenų transformacijos perkeltos į kalibratoriaus vidų. Taip pat pastebėta, kad vien daugumos klasės taisyklė yra per silpna kaštų atskaita, todėl pridėta visada bad taisyklė. Patikslinta, kad pr_auc stulpelyje skaičiuojamas AP, o fuzzy balas nėra išmokta kalibruota tikimybė. Tai konkretūs pavyzdžiai, kodėl AI pasiūlymas nelaikytas savaiminiu įrodymu.')
h('Galutinės išvados')
for txt in [
'1. Sukurtas veikiantis ir pakartojamas prototipas: nuo duomenų atsisiuntimo iki naujos paraiškos prognozės ir rezultatų peržiūros.',
'2. Pagal nustatytą 5:1 klaidų kainą pasirinktas SVM. Jo vidutinė CV kaina 0,5583 yra mažiausia tarp išbandytų variantų; statistinis pranašumas neteigtas.',
'3. Galutiniame teste 0,15 slenkstis leidžia aptikti 95 % bad atvejų, bet sukelia daug good atmetimų. Slenkstis yra svarbi sprendimo politikos dalis.',
'4. Abliacija, trūkstamų reikšmių ir grupių analizė atskleidė kompromisus. Rezultatai turi būti vertinami kartu, o ne pagal vieną geriausiai atrodantį rodiklį.',
'5. Logistinė regresija išlieka stipri paprastesnė alternatyva. Kiti medžių metodai būtų galimi tolesni bandymai; šiame darbe jie nebuvo empiriškai palyginti.']:p(txt)

page('18 Literatūra ir vertinimo kriterijų atitiktis')
refs=source_section('## Literatūra','Internetiniai šaltiniai')
for r in refs:
    q=p(r);q.paragraph_format.space_after=Pt(5)
    for run in q.runs:run.font.size=Pt(9.5)
for r in [
'[8] Breiman, L. (2001). Random Forests. Machine Learning, 45, 5–32. https://doi.org/10.1023/A:1010933404324',
'[9] Friedman, J. H. (2001). Greedy function approximation A gradient boosting machine. The Annals of Statistics, 29(5), 1189–1232. https://doi.org/10.1214/aos/1013203451']:
    q=p(r)
    for run in q.runs:run.font.size=Pt(9.5)
h('Kur ataskaitoje paaiškinti vertinimo punktai')
table(['Kriterijus','Ataskaitos skyrius'],[
('1.1 Problemos neaiškumai','1 ir 17'),('1.2 Uždavinių išskaidymas','2 ir 7'),('1.3 Prasmė įvairiems specialistams','2'),
('2.1 Alternatyvūs metodai','4 ir 5'),('2.2 Pagrindimas literatūra','6 ir 18'),('3.1 Pasirinktas sprendimas','9 ir 10'),
('3.2 Privalumai ir kompromisai','5, 9, 10 ir 13'),('4.1 Įvestis ir išvestis','3 ir 16'),
('4.2 Veiksmų seka','2, 7 ir 16'),('4.3 Formulės ir kodo ryšys','11 ir 12')],[3.9,3])
p('Eksperimento įrodymų katalogas saugykloje: results/final. Skaičių patikra: results/verification.json. Gyvo gynimo scenarijus: docs/gynimas.md. Studentas dar turi pats paaiškinti kodą, paleisti nematytą bandymą ir atlikti nedidelį pakeitimą.')

OUT.parent.mkdir(parents=True,exist_ok=True)
for par in doc.paragraphs:
    for run in par.runs:
        run.font.color.rgb=RGBColor(0,0,0)
        if par.style.name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
            run.font.name='Arial';run.font.size=doc.styles[par.style.name].font.size
for style in doc.styles:
    if style.type in (1,2):
        style.font.color.rgb=RGBColor(0,0,0)
doc.save(OUT)
print(OUT)
print('Explicit pages:',19)


